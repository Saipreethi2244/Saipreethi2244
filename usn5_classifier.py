"""
usn5_classifier.py — Input / Process (Calc) / Output classification of workbook sheets
using the USN5 "Applied Logic", with a human-readable *Classification Reason* per sheet.

Order of evaluation (mirrors the USN5 'Applied Logic' sheet):
    0. Reference lookup  – if the USN5_Data_Quantification file is supplied and the
                           (workbook, sheet) is in it, its classification/reason is reused verbatim.
    1. Mandatory PROCESS override – UTA, DTTU, Longview, CorpTax
         (sheet-name evidence first, then header/content indicators)
    2. OUTPUT – only named final-result tabs (Taxable Income Summary, TI Summary,
         Filing Package, Audit Package ...)
    3. Defined INPUT source systems – SAP, ERS, ERS - APLI, SAP Tax Module,
         Data Request from R&A, KAS, TMN02, Election Statement, GEMS Extract,
         Tax Law / Regulations  (sheet-name evidence first, then content indicators)
    4. Generic workpaper tab – Process if the tab calculates / maps / reconciles /
         allocates / adjusts / validates / combines / supports schedules, otherwise Input.

Only VISIBLE sheets are classified (hidden / veryHidden are out of scope).
The indicator lists and thresholds below were derived from the USN5 output
(they are configuration — edit the tables to tune them).
"""
from __future__ import annotations

import re
from typing import Dict, Iterable, List, Optional, Tuple

# --------------------------------------------------------------------------
# CONFIGURATION
# --------------------------------------------------------------------------
# (system, [name tokens], label used in the reason text)
PROCESS_NAME_RULES: List[Tuple[str, List[str], str]] = [
    ("UTA", ["UTA"], "UTA"),
    ("DTTU", ["DTTU"], "DTTU"),
    ("Longview", ["LV", "A9", "5100"], "LV/A9/5100"),
    ("CorpTax", ["TBBS"], "TBBS"),
]
# (system, indicators, minimum number of indicators that must match)
PROCESS_CONTENT_RULES: List[Tuple[str, List[str], int]] = [
    ("Longview", ["LVL4", "TAXCOLUMNS", "TIMEPER", "5100"], 1),
    ("CorpTax", ["TBBS", "COMBINED BALANCE SHEET"], 1),
]
OUTPUT_NAME_PHRASES = ["TAXABLE INCOME SUMMARY", "TI SUMMARY", "FILING PACKAGE", "AUDIT PACKAGE"]

INPUT_NAME_RULES: List[Tuple[str, List[str], str]] = [
    ("ERS - APLI", ["APLI"], "APLI"),
    ("ERS", ["ERS"], "ERS"),
    ("SAP", ["SAP"], "SAP"),
    ("Data Request from R&A", ["A202"], "A202"),
    ("KAS", ["KAS"], "KAS"),
    ("TMN02", ["TMN02"], "TMN02"),
    ("Election Statement", ["ELECTION"], "ELECTION"),
    ("GEMS Extract", ["GEMS"], "GEMS"),
    ("Tax Law / Regulations", ["REGULATIONS", "REGS"], "REGULATIONS"),
]
# ordered: on equal match count the earlier system wins
INPUT_CONTENT_RULES: List[Tuple[str, List[str], int]] = [
    ("SAP Tax Module", ["DEPRECIATION KEY", "SUB-NUMBER", "ASSET DESCRIPTION", "ASSET CLASS"], 2),
    ("ERS - APLI", ["DOC # FI", "DOC TYPE FI", "LE:OU LEGAL ENTITY", "LINE ITEM TEXT ORIG", "POSTING PERIOD"], 3),
    ("ERS", ["G/L ACCOUNT KEY", "ITD BALANCE AMOUNT"], 2),
    ("SAP", ["COMPANY CODE", "ACCOUNT NUMBER", "ACCOUNT NUMBER DESCRIPTION", "POSTING LEGAL ENTITY", "LOCAL JV CODE"], 2),
    ("TMN02", ["TMN02"], 1),
]
# indicator order matches the wording used in the USN5 reasons
CALC_INDICATORS = ["CALCULATION", "RECONCILIATION", "MAPPING", "ANALYSIS", "ROLLFORWARD",
                   "ADJUSTMENT", "WORKPAPER", "SUMMARY", "SCHEDULE"]
SUPPORT_NAME_HINTS = ["NOTE", "SUPPORT", "MEMO", "COMMENT"]

_TOKEN = re.compile(r"[A-Z0-9]+")


def _name_tokens(name: str) -> List[str]:
    return _TOKEN.findall(name.upper())


def _has(text: str, ind: str) -> bool:
    return re.search(r"(?<![A-Z0-9])" + re.escape(ind) + r"(?![A-Z0-9])", text) is not None


def _hits(header_blob: str, inds: List[str]) -> List[str]:
    return [i for i in inds if _has(header_blob, i)]


# --------------------------------------------------------------------------
# core rule engine
# --------------------------------------------------------------------------
def classify_sheet(name: str, n_formula: int, n_nonblank: int,
                   header_text: Iterable[str]) -> Tuple[str, str, str]:
    """Return (classification, source_system, reason) for one visible sheet."""
    up = name.upper()
    tokens = _name_tokens(name)
    blob = " | ".join(header_text)

    # 1a. mandatory PROCESS override — sheet name
    for system, toks, label in PROCESS_NAME_RULES:
        hit = next((t for t in toks if t in tokens), None)
        if hit:
            lab = label if len(toks) > 2 else hit
            return "Process", system, f"sheet name contains {lab}; mandatory process override for {system}"
    # 1b. mandatory PROCESS override — content
    for system, inds, need in PROCESS_CONTENT_RULES:
        h = _hits(blob, inds)
        if len(h) >= need:
            return "Process", system, f"content indicators: {', '.join(h)}; mandatory process override for {system}"

    # 2. OUTPUT — named deliverable tabs only
    for ph in OUTPUT_NAME_PHRASES:
        if ph in up or ph.replace(" ", "_") in up:
            return "Output", "Workbook / Tax workpaper", \
                f"sheet name '{name.strip()}' is a named final-result/deliverable tab ({ph.title()})"

    # 3a. defined INPUT systems — sheet name
    for system, toks, label in INPUT_NAME_RULES:
        if any(t in tokens for t in toks):
            return "Input", system, f"sheet name contains {label}; {system} is defined as input"
    # 3b. defined INPUT systems — content
    best = None
    for system, inds, need in INPUT_CONTENT_RULES:
        h = _hits(blob, inds)
        if len(h) >= need and (best is None or len(h) > len(best[1])):
            best = (system, h)
    if best:
        return "Input", best[0], f"content indicators: {', '.join(best[1])}; {best[0]} is defined as input"

    # 4. generic workpaper tab
    base = ("no source system met the required matching-indicator threshold; "
            "sheet functions within the workpaper; ")
    ind_hits = [i for i in CALC_INDICATORS if _has(up + " | " + blob, i)]
    ind_txt = f"; indicators: {', '.join(ind_hits)}" if ind_hits else ""
    if n_formula > 0:
        return "Process", "Workbook / Tax workpaper", base + f"performs work through {n_formula} formula cells{ind_txt}"
    if ind_hits:
        return "Process", "Workbook / Tax workpaper", base + f"performs work through 0 formula cells{ind_txt}"
    if n_nonblank > 0 and any(h in up for h in SUPPORT_NAME_HINTS):
        return "Process", "Workbook / Tax workpaper", base + "workpaper support tab; not a named final deliverable"
    return "Input", "Workbook / Tax workpaper", \
        base + "visible data appears to be source data with no calculation/reconciliation indicators"


# --------------------------------------------------------------------------
# USN5 reference file (optional) — reuse the already-approved classification
# --------------------------------------------------------------------------
_YEAR = re.compile(r"(?<!\d)(?:19|20)\d{2}(?!\d)")


def _fkey(fname: str) -> str:
    return _YEAR.sub("#", fname.strip().lower())


def _skey(sname: str) -> str:
    return re.sub(r"\s+", " ", sname.strip().lower())


def load_usn5_reference(xlsx_bytes: bytes) -> Dict[Tuple[str, str], dict]:
    """Read the 'Data Quantification' sheet of a USN5 file -> {(file, sheet): row}.
    Uses read_only openpyxl (≈ 0.1 s for 226 rows)."""
    import io
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    ws = wb["Data Quantification"] if "Data Quantification" in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    head = [str(h).strip() if h is not None else "" for h in next(it)]
    idx = {h: i for i, h in enumerate(head)}
    ref: Dict[Tuple[str, str], dict] = {}
    for r in it:
        if not r or r[idx["Sheet Name"]] is None:
            continue
        row = {h: r[i] for h, i in idx.items()}
        ref[(_fkey(str(row["File Name"])), _skey(str(row["Sheet Name"])))] = row
    wb.close()
    return ref


def classify_workbook(file_name: str, scans: dict, reference: Optional[Dict[Tuple[str, str], dict]] = None
                      ) -> List[dict]:
    """scans = {sheet: SheetScan} (from formula_flow.scan_workbook).  Returns one row per visible sheet."""
    rows = []
    fk = _fkey(file_name)
    for name, sc in scans.items():
        ref = reference.get((fk, _skey(name))) if reference else None
        if ref:
            rows.append({
                "File Name": file_name, "Sheet Name": name,
                "Classification": ref["Classification"], "Source System": ref.get("Source System"),
                "Adjustment Category": ref.get("Adjustment Category"),
                "Classification Reason": ref.get("Classification Reason"),
                "Basis": "USN5 reference file",
                "Visible Used Rows": sc.used_rows, "Visible Used Columns": sc.used_cols,
                "Nonblank Cells": sc.nonblank, "Formula Cells": sc.n_formula,
                "Manual (numeric) Cells": sc.n_manual,
            })
            continue
        cls, system, reason = classify_sheet(name, sc.n_formula, sc.nonblank, sc.header_text)
        rows.append({
            "File Name": file_name, "Sheet Name": name, "Classification": cls,
            "Source System": system, "Adjustment Category": None,
            "Classification Reason": reason, "Basis": "Rule engine",
            "Visible Used Rows": sc.used_rows, "Visible Used Columns": sc.used_cols,
            "Nonblank Cells": sc.nonblank, "Formula Cells": sc.n_formula,
            "Manual (numeric) Cells": sc.n_manual,
        })
    return rows
